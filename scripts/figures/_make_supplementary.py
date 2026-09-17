"""Build the supplementary tables (XLSX) for the Functional & Integrative Genomics submission.

Everything is derived from the existing result files; no number is typed in by hand.

  ESM_1  per-cancer network summary (33 rows)                = Online Resource 1 / Table S1
  ESM_2  all inferred causal edges with cross-cancer sharing  = Online Resource 2 / Table S2
  ESM_3  composition of the shared set (>=3 cancers) + DepMap + hub specificity
  ESM_4  MSigDB C2 pathway enrichment, per-cancer and pan-cancer

The journal asks for the article title, journal name, author names and the affiliation and
e-mail address of the corresponding author to be included in every supplementary FILE, so each
sheet opens with that identification block before the data header.
"""
import os, sys, json
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8', errors='replace')
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _gene_symbols import canonical  # noqa: E402  (same directory; keeps table symbols clean)

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
RES = os.path.join(BASE, 'results')
OUT = os.path.join(BASE, 'supplementary')
os.makedirs(OUT, exist_ok=True)

HDR = Font(bold=True)
WRAP = Alignment(vertical='top', wrap_text=False)

# Required in every supplementary file by the journal's submission guidelines.
BANNER = [
    'Article title: Edge-level causal graph comparison across 33 TCGA cohorts reveals '
    'tissue-specific regulatory structure and prognostic hub genes',
    'Journal: Functional & Integrative Genomics',
    'Author: Shuaidong Gao',
    'Affiliation: Chongqing Institute of Foreign Studies, Qijiang Campus, '
    'Chongqing 401420, China',
    'Corresponding author e-mail: gaoshuaidong@stu.cqifs.edu.cn',
]


def write_sheet(ws, header, rows, widths=None):
    for line in BANNER:
        ws.append([line])
    ws.append([''])
    ws.append(header)
    hrow = ws.max_row
    for c in range(1, len(header) + 1):
        ws.cell(row=hrow, column=c).font = HDR
    for r in rows:
        ws.append(list(r))
    ws.freeze_panes = 'A%d' % (hrow + 1)
    for i, w in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws


# ---------------------------------------------------------------- inputs
nt = json.load(open(os.path.join(RES, '_pipeline_notears.json'), encoding='utf-8'))
pc = json.load(open(os.path.join(RES, '_per_cancer_network.json'), encoding='utf-8'))
share = json.load(open(os.path.join(RES, '_cancer_pair_sharing.json'), encoding='utf-8'))
dep = json.load(open(os.path.join(RES, '_depmap_validation.json'), encoding='utf-8'))
tau = json.load(open(os.path.join(RES, '_tau_specificity.json'), encoding='utf-8'))
hub = json.load(open(os.path.join(RES, '_hub_expression.json'), encoding='utf-8'))
per_enr = json.load(open(os.path.join(RES, '_enrichment_percancer.json'), encoding='utf-8'))['per_cancer']
glo_enr = json.load(open(os.path.join(RES, '_enrichment_global.json'), encoding='utf-8'))

cancers = sorted(pc.keys())
TAU = 0.3

# ---------------------------------------------------------------- S1
wb = Workbook()
ws = wb.active
ws.title = 'per-cancer summary'
rows = []
for c in cancers:
    W = np.array(nt[c]['W'], dtype=float)
    A = (np.abs(W) > TAU).astype(float)
    np.fill_diagonal(A, 0.0)
    deg = A.sum(axis=0) + A.sum(axis=1)
    nz = np.abs(W[np.abs(W) > TAU])
    hg = pc[c]['hub']
    t_hub = tau['gene_tau'].get(hg)
    r_hub = hub['stats'].get(hg, {}).get('rank')
    rows.append([c, int(nt[c].get('n', 0)), int(pc[c]['n_edges']),
                 round(float(np.mean(nz)), 4), round(float(np.max(nz)), 4),
                 hg, int(pc[c]['hub_degree']),
                 round(float(t_hub), 4) if t_hub is not None else 'not assessed',
                 int(r_hub) if r_hub is not None else 'not assessed',
                 round(float(deg.max() / max(deg.sum(), 1)), 4)])
write_sheet(ws, ['Cancer', 'Samples (n)', 'Causal edges (|w|>0.3)', 'Mean |w|', 'Max |w|',
                 'Hub gene (max degree)', 'Hub degree', 'Hub tau', 'Hub own-cancer rank',
                 'Hub degree share'], rows,
            [12, 12, 20, 10, 10, 22, 12, 10, 18, 16])
p1 = os.path.join(OUT, 'ESM_1.xlsx')
wb.save(p1)
print('wrote %s  (%d rows)' % (os.path.basename(p1), len(rows)))

# ---------------------------------------------------------------- S2
pair_count = {}
for c in cancers:
    W = np.array(nt[c]['W'], dtype=float)
    genes = nt[c]['genes']
    for i in range(100):
        for j in range(100):
            if i != j and abs(W[i, j]) > TAU:
                pair_count.setdefault((genes[i], genes[j]), set()).add(c)

wb = Workbook()
ws = wb.active
ws.title = 'all causal edges'
rows = []
for c in cancers:
    W = np.array(nt[c]['W'], dtype=float)
    genes = nt[c]['genes']
    for i in range(100):
        for j in range(100):
            if i != j and abs(W[i, j]) > TAU:
                w_ = float(W[i, j])
                others = sorted(pair_count[(genes[i], genes[j])] - {c})
                # Look the pair up under the spelling the matrix carries, print it under the
                # HGNC-approved one: three rows of HiSeqV2 have no symbol and arrive as "?|<entrez>".
                rows.append([c, canonical(genes[i]), canonical(genes[j]), round(w_, 5),
                             'positive' if w_ > 0 else 'negative',
                             len(others) + 1, ';'.join(others)])
write_sheet(ws, ['Cancer', 'Source gene', 'Target gene', 'Weight w', 'Sign',
                 'Detected in N cancers', 'Also detected in'], rows,
            [10, 16, 16, 12, 10, 20, 40])
p2 = os.path.join(OUT, 'ESM_2.xlsx')
wb.save(p2)
print('wrote %s  (%d rows)' % (os.path.basename(p2), len(rows)))

# ---------------------------------------------------------------- S3
wb = Workbook()
ws = wb.active
ws.title = 'shared >=3 cancers'
comp = share['shared_ge3_pairs']
rows = []
for kind, pairs in [('sex-chromosome / X-inactivation', comp['sex_chromosome']),
                    ('same gene family', comp['same_family']),
                    ('other', comp['other'])]:
    for a, b in pairs:
        n = pair_count.get((a, b), pair_count.get((b, a), set()))
        rows.append([a, b, kind, len(n) if n else '', ';'.join(sorted(n)) if n else ''])
write_sheet(ws, ['Gene A', 'Gene B', 'Classification', 'Detected in N cancers', 'Cancers'],
            rows, [16, 16, 34, 20, 26])

ws2 = wb.create_sheet('DepMap concordance')
rows2 = [[r['A'], r['B'], ';'.join(r['cancers']), r['n_celllines'],
          round(r['expr_corr'], 4), '%.3g' % r['expr_p'],
          round(r['crisprA_median'], 4), round(r['crisprB_median'], 4),
          '' if r['crispr_corr'] != r['crispr_corr'] else round(r['crispr_corr'], 4),
          '' if r['crispr_p'] != r['crispr_p'] else '%.3g' % r['crispr_p']]
         for r in sorted(dep, key=lambda x: -x['expr_corr'])]
write_sheet(ws2, ['Gene A', 'Gene B', 'TCGA cancers', 'Cell lines (n)',
                  'Expression Spearman r', 'Expression p',
                  'CRISPR effect median A', 'CRISPR effect median B',
                  'CRISPR co-dependency r', 'CRISPR p'],
            rows2, [10, 10, 18, 14, 20, 14, 20, 20, 20, 12])

ws3 = wb.create_sheet('hub tissue specificity')
rows3 = []
for g in sorted(hub['stats'], key=lambda x: hub['stats'][x]['rank']):
    v = hub['stats'][g]
    t_g = tau['gene_tau'].get(g)
    rows3.append([g, v['own_cancer'], v['rank'], v['n_cancers_with_gene'],
                  round(v['own_median'], 3), round(v['max_median'], 3), v['max_cancer'],
                  ';'.join(v['top5']),
                  round(float(t_g), 4) if t_g is not None else 'not assessed'])
write_sheet(ws3, ['Hub gene', 'Own cancer', 'Rank of own cancer (of 33)',
                  'Cohorts with expression record', 'Own-cancer median (log2 TPM+1)',
                  'Max median', 'Cancer with max median', 'Top-5 cohorts by median',
                  'Tissue-specificity index tau'], rows3,
            [12, 12, 24, 28, 28, 12, 22, 34, 26])
p3 = os.path.join(OUT, 'ESM_3.xlsx')
wb.save(p3)
print('wrote %s  (%d + %d + %d rows)' % (os.path.basename(p3), len(rows), len(rows2), len(rows3)))

# ---------------------------------------------------------------- S4
wb = Workbook()
ws = wb.active
ws.title = 'per-cancer enrichment'
rows = []
for c in sorted(per_enr.keys()):
    b = per_enr[c]
    for e in sorted(b['enrichment'], key=lambda x: x['p']):
        rows.append([c, b['n_test_genes'], b['n_universe'], e['set'], e['k'], e['K'],
                     round(float(e['enrichment']), 4), '%.4g' % e['p'],
                     'yes' if e['p'] < 0.05 else 'no'])
write_sheet(ws, ['Cancer', 'Network genes (n)', 'Universe (n)', 'MSigDB C2 gene set',
                 'k (hits)', 'K (set size)', 'Enrichment fold', 'p-value', 'p < 0.05'],
            rows, [10, 18, 14, 56, 10, 12, 16, 12, 10])

ws2 = wb.create_sheet('pan-cancer enrichment')
rows2 = [[e['set'], e['k'], e['K'], glo_enr['global_network_genes'], glo_enr['universe_size'],
          round(float(e['enrichment']), 4), '%.4g' % e['p']]
         for e in sorted(glo_enr['global_enrichment'], key=lambda x: x['p'])]
write_sheet(ws2, ['MSigDB C2 gene set', 'k (hits)', 'K (set size)', 'Network genes (n)',
                  'Universe (n)', 'Enrichment fold', 'p-value'],
            rows2, [56, 10, 12, 18, 14, 16, 12])
p4 = os.path.join(OUT, 'ESM_4.xlsx')
wb.save(p4)
print('wrote %s  (%d + %d rows)' % (os.path.basename(p4), len(rows), len(rows2)))
print('DONE')
